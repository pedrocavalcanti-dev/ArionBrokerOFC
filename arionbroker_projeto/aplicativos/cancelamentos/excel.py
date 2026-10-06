"""
excel.py — Relatório profissional do módulo de Cancelamentos.

Responsabilidades:
- filtrar e consultar os cancelamentos;
- gerar workbook XLSX;
- montar Resumo, Cancelamentos e Pendências;
- aplicar formatação profissional;
- gerar filtros válidos sem Table/AutoFilter duplicados;
- controlar bordas, larguras, alinhamentos e alturas;
- gerar gráfico de colunas do Resumo;
- devolver o arquivo pronto para download.

IMPORTANTE:
Não usamos openpyxl.worksheet.table.Table neste relatório.
Isso evita o conflito de Table + AutoFilter que fazia o Microsoft Excel
reparar/remover partes do arquivo gerado.
"""

from io import BytesIO
from datetime import datetime

from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
from django.utils import timezone
from django.views.decorators.http import require_http_methods

from .models import Cancelamento, PendenciaCancelamento

try:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    Workbook = None


# =============================================================================
# PALETA
# =============================================================================
COR_NAVY = '0F172A'
COR_NAVY_2 = '172033'
COR_VERDE = '16A34A'
COR_VERDE_CLARO = 'DCFCE7'
COR_AZUL = '2563EB'
COR_AZUL_CLARO = 'DBEAFE'
COR_ROXO = '7C3AED'
COR_ROXO_CLARO = 'EDE9FE'
COR_AMARELO = 'D97706'
COR_AMARELO_CLARO = 'FEF3C7'
COR_VERMELHO = 'DC2626'
COR_VERMELHO_CLARO = 'FEE2E2'
COR_CINZA = '64748B'
COR_CINZA_2 = '94A3B8'
COR_CINZA_CLARO = 'F8FAFC'
COR_CINZA_ZEBRA = 'F1F5F9'
COR_BORDA = 'CBD5E1'
COR_BRANCO = 'FFFFFF'
COR_PRETO = '111827'


# =============================================================================
# HELPERS DE DADOS
# =============================================================================
def _nome_usuario(usuario, vazio='Sem responsável'):
    if not usuario:
        return vazio

    return (
        usuario.get_full_name()
        or usuario.username
        or vazio
    )


def _numero(valor):
    if valor in (None, ''):
        return 0.0

    try:
        return float(valor)
    except (TypeError, ValueError):
        return 0.0


def _data_excel(valor):
    if not valor:
        return None

    if isinstance(valor, datetime):
        try:
            if timezone.is_aware(valor):
                valor = timezone.localtime(valor)
        except Exception:
            pass

        return valor.replace(tzinfo=None)

    return valor


def _pendencias_validas(cancelamento):
    return [
        p
        for p in cancelamento.pendencias.all()
        if (
            p.status
            != PendenciaCancelamento.StatusPendencia.DESCARTADA
        )
    ]


# =============================================================================
# HELPERS DE ESTILO
# =============================================================================
BORDA_FINA = Side(
    style='thin',
    color=COR_BORDA,
)

BORDA_TABELA = Border(
    left=BORDA_FINA,
    right=BORDA_FINA,
    top=BORDA_FINA,
    bottom=BORDA_FINA,
)


def _fill(cor):
    return PatternFill(
        fill_type='solid',
        fgColor=cor,
    )


def _titulo_planilha(
    ws,
    titulo,
    subtitulo,
    ultima_coluna,
):
    ws.merge_cells(
        start_row=1,
        start_column=1,
        end_row=1,
        end_column=ultima_coluna,
    )

    celula = ws.cell(
        row=1,
        column=1,
        value=titulo,
    )

    celula.fill = _fill(COR_NAVY)
    celula.font = Font(
        color=COR_BRANCO,
        bold=True,
        size=17,
    )
    celula.alignment = Alignment(
        horizontal='center',
        vertical='center',
    )
    celula.border = Border(
        bottom=Side(
            style='medium',
            color=COR_VERDE,
        )
    )

    ws.row_dimensions[1].height = 32

    if subtitulo:
        ws.merge_cells(
            start_row=2,
            start_column=1,
            end_row=2,
            end_column=ultima_coluna,
        )

        sub = ws.cell(
            row=2,
            column=1,
            value=subtitulo,
        )

        sub.font = Font(
            color=COR_CINZA,
            italic=True,
            size=10,
        )
        sub.alignment = Alignment(
            horizontal='center',
            vertical='center',
        )

        ws.row_dimensions[2].height = 21


def _cabecalho(
    ws,
    linha,
    col_inicio,
    col_fim,
):
    for coluna in range(
        col_inicio,
        col_fim + 1,
    ):
        celula = ws.cell(
            row=linha,
            column=coluna,
        )

        celula.fill = _fill(COR_NAVY)
        celula.font = Font(
            color=COR_BRANCO,
            bold=True,
            size=9,
        )
        celula.alignment = Alignment(
            horizontal='center',
            vertical='center',
            wrap_text=True,
        )
        celula.border = BORDA_TABELA

    ws.row_dimensions[linha].height = 31


def _grupo_colunas(
    ws,
    linha,
    inicio,
    fim,
    titulo,
    cor,
):
    # Escreve e estiliza a célula principal antes do merge.
    # Isso evita AttributeError em MergedCell.
    c = ws.cell(
        row=linha,
        column=inicio,
        value=titulo,
    )

    c.fill = _fill(cor)
    c.font = Font(
        color=COR_BRANCO,
        bold=True,
        size=9,
    )
    c.alignment = Alignment(
        horizontal='center',
        vertical='center',
    )
    c.border = BORDA_TABELA

    if fim > inicio:
        ws.merge_cells(
            start_row=linha,
            start_column=inicio,
            end_row=linha,
            end_column=fim,
        )

    ws.row_dimensions[linha].height = 21


def _estilizar_corpo(
    ws,
    inicio,
    fim,
    total_colunas,
    wrap_cols=None,
    center_cols=None,
):
    wrap_cols = set(
        wrap_cols or []
    )

    center_cols = set(
        center_cols or []
    )

    for linha in range(
        inicio,
        fim + 1,
    ):
        cor_linha = (
            COR_BRANCO
            if (linha - inicio) % 2 == 0
            else COR_CINZA_ZEBRA
        )

        for coluna in range(
            1,
            total_colunas + 1,
        ):
            celula = ws.cell(
                row=linha,
                column=coluna,
            )

            celula.fill = _fill(
                cor_linha
            )

            celula.border = BORDA_TABELA

            celula.font = Font(
                color=COR_PRETO,
                size=9,
            )

            celula.alignment = Alignment(
                vertical='top',
                horizontal=(
                    'center'
                    if coluna in center_cols
                    else 'left'
                ),
                wrap_text=(
                    coluna in wrap_cols
                ),
            )

        # Mantém legível sem transformar observação longa
        # em uma linha gigantesca.
        ws.row_dimensions[linha].height = 30


def _larguras(
    ws,
    mapa,
):
    for coluna, largura in mapa.items():
        ws.column_dimensions[
            coluna
        ].width = largura


def _numero_moeda(
    ws,
    colunas,
    linha_inicio,
    linha_fim,
):
    for linha in range(
        linha_inicio,
        linha_fim + 1,
    ):
        for coluna in colunas:
            ws[
                f'{coluna}{linha}'
            ].number_format = (
                'R$ #,##0.00;[Red]-R$ #,##0.00'
            )


def _numero_data(
    ws,
    colunas,
    linha_inicio,
    linha_fim,
    com_hora=False,
):
    formato = (
        'dd/mm/yyyy hh:mm'
        if com_hora
        else 'dd/mm/yyyy'
    )

    for linha in range(
        linha_inicio,
        linha_fim + 1,
    ):
        for coluna in colunas:
            ws[
                f'{coluna}{linha}'
            ].number_format = formato


def _sem_grade(ws):
    ws.sheet_view.showGridLines = False


# =============================================================================
# FILTROS DA EXPORTAÇÃO
# =============================================================================
def _queryset_relatorio(request):
    qs = (
        Cancelamento.objects
        .select_related(
            'responsavel',
            'atendente',
        )
        .prefetch_related(
            'observacoes',
            'pendencias__criada_por',
        )
    )

    filas_validas = {
        valor
        for valor, _label
        in Cancelamento.Fila.choices
    }

    filas = [
        fila
        for fila
        in request.GET.getlist('fila')
        if fila in filas_validas
    ]

    if filas:
        qs = qs.filter(
            fila__in=filas
        )

    responsavel = request.GET.get(
        'responsavel',
        '',
    ).strip()

    if responsavel == 'meus':
        qs = qs.filter(
            responsavel=request.user
        )

    elif responsavel == 'sem_responsavel':
        qs = qs.filter(
            responsavel__isnull=True
        )

    elif responsavel.isdigit():
        qs = qs.filter(
            responsavel_id=int(responsavel)
        )

    status = request.GET.get(
        'status',
        '',
    ).strip()

    if status:
        qs = qs.filter(
            status=status
        )

    data_inicio = request.GET.get(
        'data_inicio',
        '',
    ).strip()

    if data_inicio:
        try:
            inicio = datetime.strptime(
                data_inicio,
                '%Y-%m-%d',
            ).date()

            qs = qs.filter(
                data_criacao_chamado__gte=inicio
            )

        except ValueError:
            pass

    data_fim = request.GET.get(
        'data_fim',
        '',
    ).strip()

    if data_fim:
        try:
            fim = datetime.strptime(
                data_fim,
                '%Y-%m-%d',
            ).date()

            qs = qs.filter(
                data_criacao_chamado__lte=fim
            )

        except ValueError:
            pass

    return qs.order_by(
        '-atualizado_em',
        '-id',
    )


# =============================================================================
# ABA RESUMO
# =============================================================================
def _montar_resumo(
    wb,
    cancelamentos,
    usuario,
):
    ws = wb.active
    ws.title = 'Resumo'

    _sem_grade(ws)

    _titulo_planilha(
        ws,
        'Relatório de Cancelamentos — ARION Broker',
        (
            f'Gerado em '
            f'{timezone.localtime():%d/%m/%Y %H:%M} '
            f'por {_nome_usuario(usuario, "Usuário")}'
        ),
        10,
    )

    total = len(
        cancelamentos
    )

    filas = {
        valor: 0
        for valor, _label
        in Cancelamento.Fila.choices
    }

    por_responsavel = {}

    pendencias_ativas = 0
    valor_pendencias = 0.0
    sem_responsavel = 0

    for c in cancelamentos:
        filas[c.fila] = (
            filas.get(
                c.fila,
                0,
            )
            + 1
        )

        if not c.responsavel_id:
            sem_responsavel += 1

        nome = _nome_usuario(
            c.responsavel
        )

        por_responsavel.setdefault(
            nome,
            {
                'total': 0,
                'pedido': 0,
                'solicitacao': 0,
                'pendencia': 0,
                'concluido': 0,
            },
        )

        dados = por_responsavel[
            nome
        ]

        dados['total'] += 1

        if (
            c.fila
            == Cancelamento.Fila.PEDIDO
        ):
            dados['pedido'] += 1

        elif (
            c.fila
            == Cancelamento.Fila.SOLICITACAO
        ):
            dados['solicitacao'] += 1

        elif (
            c.fila
            == Cancelamento.Fila.PENDENCIA
        ):
            dados['pendencia'] += 1

        elif (
            c.fila
            == Cancelamento.Fila.CONCLUIDO
        ):
            dados['concluido'] += 1

        validas = _pendencias_validas(
            c
        )

        pendencias_ativas += len(
            validas
        )

        valor_pendencias += sum(
            _numero(
                p.valor
            )
            for p in validas
        )

    # -------------------------------------------------------------------------
    # KPI CARDS
    # -------------------------------------------------------------------------
    kpis = [
        (
            'Total Geral',
            total,
            COR_AZUL_CLARO,
            COR_AZUL,
            False,
        ),
        (
            'Pedido',
            filas.get(
                Cancelamento.Fila.PEDIDO,
                0,
            ),
            COR_ROXO_CLARO,
            COR_ROXO,
            False,
        ),
        (
            'Solicitação',
            filas.get(
                Cancelamento.Fila.SOLICITACAO,
                0,
            ),
            COR_AZUL_CLARO,
            COR_AZUL,
            False,
        ),
        (
            'Pendências',
            filas.get(
                Cancelamento.Fila.PENDENCIA,
                0,
            ),
            COR_AMARELO_CLARO,
            COR_AMARELO,
            False,
        ),
        (
            'Concluídos',
            filas.get(
                Cancelamento.Fila.CONCLUIDO,
                0,
            ),
            COR_VERDE_CLARO,
            COR_VERDE,
            False,
        ),
        (
            'Sem responsável',
            sem_responsavel,
            COR_VERMELHO_CLARO,
            COR_VERMELHO,
            False,
        ),
        (
            'Pendências ativas',
            pendencias_ativas,
            COR_AMARELO_CLARO,
            COR_AMARELO,
            False,
        ),
        (
            'Valor das pendências',
            valor_pendencias,
            COR_VERDE_CLARO,
            COR_VERDE,
            True,
        ),
    ]

    posicoes = [
        (4, 1, 2),
        (4, 3, 4),
        (4, 5, 6),
        (4, 7, 8),
        (7, 1, 2),
        (7, 3, 4),
        (7, 5, 6),
        (7, 7, 8),
    ]

    for (
        label,
        valor,
        fundo,
        destaque,
        moeda,
    ), (
        linha,
        col_ini,
        col_fim,
    ) in zip(
        kpis,
        posicoes,
    ):
        ws.merge_cells(
            start_row=linha,
            start_column=col_ini,
            end_row=linha,
            end_column=col_fim,
        )

        ws.merge_cells(
            start_row=linha + 1,
            start_column=col_ini,
            end_row=linha + 1,
            end_column=col_fim,
        )

        label_cell = ws.cell(
            row=linha,
            column=col_ini,
            value=label,
        )

        valor_cell = ws.cell(
            row=linha + 1,
            column=col_ini,
            value=valor,
        )

        for l in (
            linha,
            linha + 1,
        ):
            for col in range(
                col_ini,
                col_fim + 1,
            ):
                cell = ws.cell(
                    row=l,
                    column=col,
                )

                cell.fill = _fill(
                    fundo
                )

                cell.border = BORDA_TABELA

        label_cell.font = Font(
            color=COR_CINZA,
            bold=True,
            size=9,
        )

        label_cell.alignment = Alignment(
            horizontal='center',
            vertical='center',
        )

        valor_cell.font = Font(
            color=destaque,
            bold=True,
            size=18,
        )

        valor_cell.alignment = Alignment(
            horizontal='center',
            vertical='center',
        )

        if moeda:
            valor_cell.number_format = (
                'R$ #,##0.00'
            )

        ws.row_dimensions[
            linha
        ].height = 22

        ws.row_dimensions[
            linha + 1
        ].height = 31

    # -------------------------------------------------------------------------
    # TABELA POR FILA
    # -------------------------------------------------------------------------
    linha_titulo_fila = 11

    ws.merge_cells(
        start_row=linha_titulo_fila,
        start_column=1,
        end_row=linha_titulo_fila,
        end_column=3,
    )

    c = ws.cell(
        linha_titulo_fila,
        1,
        'Distribuição por Fila',
    )

    c.fill = _fill(
        COR_NAVY
    )

    c.font = Font(
        color=COR_BRANCO,
        bold=True,
        size=11,
    )

    c.alignment = Alignment(
        horizontal='center',
        vertical='center',
    )

    for col in range(
        1,
        4,
    ):
        ws.cell(
            linha_titulo_fila,
            col,
        ).border = BORDA_TABELA

    headers_fila = [
        'Fila',
        'Quantidade',
        '% do Total',
    ]

    linha_header_fila = (
        linha_titulo_fila + 1
    )

    for col, header in enumerate(
        headers_fila,
        start=1,
    ):
        ws.cell(
            linha_header_fila,
            col,
            header,
        )

    _cabecalho(
        ws,
        linha_header_fila,
        1,
        3,
    )

    linhas_fila = [
        (
            'Pedido',
            filas.get(
                Cancelamento.Fila.PEDIDO,
                0,
            ),
        ),
        (
            'Solicitação',
            filas.get(
                Cancelamento.Fila.SOLICITACAO,
                0,
            ),
        ),
        (
            'Pendências',
            filas.get(
                Cancelamento.Fila.PENDENCIA,
                0,
            ),
        ),
        (
            'Concluídos',
            filas.get(
                Cancelamento.Fila.CONCLUIDO,
                0,
            ),
        ),
    ]

    row = linha_header_fila + 1

    for label, quantidade in linhas_fila:
        ws.cell(
            row,
            1,
            label,
        )

        ws.cell(
            row,
            2,
            quantidade,
        )

        ws.cell(
            row,
            3,
            (
                quantidade / total
                if total
                else 0
            ),
        )

        row += 1

    fim_fila = row - 1

    _estilizar_corpo(
        ws,
        linha_header_fila + 1,
        fim_fila,
        3,
        center_cols={
            2,
            3,
        },
    )

    for r in range(
        linha_header_fila + 1,
        fim_fila + 1,
    ):
        ws.cell(
            r,
            3,
        ).number_format = '0.0%'

    # -------------------------------------------------------------------------
    # GRÁFICO PROFISSIONAL — COLUNAS
    # -------------------------------------------------------------------------
    if total > 0:
        chart = BarChart()
        chart.type = 'col'
        chart.style = 10
        chart.title = 'Chamados por Fila'
        chart.height = 8.2
        chart.width = 14.2
        chart.legend = None

        # Eixo vertical mostra apenas a quantidade.
        chart.y_axis.title = 'Quantidade'
        chart.y_axis.scaling.min = 0
        chart.y_axis.majorGridlines = None

        # Eixo horizontal usa os nomes das filas.
        chart.x_axis.title = None
        chart.x_axis.tickLblPos = 'low'

        data = Reference(
            ws,
            min_col=2,
            min_row=linha_header_fila,
            max_row=fim_fila,
        )

        categorias = Reference(
            ws,
            min_col=1,
            min_row=linha_header_fila + 1,
            max_row=fim_fila,
        )

        chart.add_data(
            data,
            titles_from_data=True,
        )

        chart.set_categories(
            categorias
        )

        # Sem DataLabels e sem legenda:
        # evita textos sobrepostos como
        # "Quantidade; Solicitação; 29".
        chart.dataLabels = None

        # Barras um pouco mais largas e legíveis.
        chart.gapWidth = 65

        ws.add_chart(
            chart,
            'E11',
        )

    # -------------------------------------------------------------------------
    # RESPONSÁVEIS
    # -------------------------------------------------------------------------
    linha_resp = 20

    ws.merge_cells(
        start_row=linha_resp,
        start_column=1,
        end_row=linha_resp,
        end_column=6,
    )

    t = ws.cell(
        linha_resp,
        1,
        'Chamados por Responsável',
    )

    t.fill = _fill(
        COR_NAVY
    )

    t.font = Font(
        color=COR_BRANCO,
        bold=True,
        size=11,
    )
    t.alignment = Alignment(
        horizontal='center',
        vertical='center',
    )

    for col in range(
        1,
        7,
    ):
        ws.cell(
            linha_resp,
            col,
        ).border = BORDA_TABELA

    headers_resp = [
        'Responsável',
        'Total',
        'Pedido',
        'Solicitação',
        'Pendências',
        'Concluídos',
    ]

    header_resp = linha_resp + 1

    for col, header in enumerate(
        headers_resp,
        start=1,
    ):
        ws.cell(
            header_resp,
            col,
            header,
        )

    _cabecalho(
        ws,
        header_resp,
        1,
        6,
    )

    row = header_resp + 1

    ordenados = sorted(
        por_responsavel.items(),
        key=lambda item: (
            -item[1]['total'],
            item[0].casefold(),
        ),
    )

    for nome, dados in ordenados:
        valores = [
            nome,
            dados['total'],
            dados['pedido'],
            dados['solicitacao'],
            dados['pendencia'],
            dados['concluido'],
        ]

        for col, valor in enumerate(
            valores,
            start=1,
        ):
            ws.cell(
                row,
                col,
                valor,
            )

        row += 1

    fim_resp = row - 1

    if fim_resp >= header_resp + 1:
        _estilizar_corpo(
            ws,
            header_resp + 1,
            fim_resp,
            6,
            center_cols={
                2,
                3,
                4,
                5,
                6,
            },
        )

        # Filtro simples e válido.
        ws.auto_filter.ref = (
            f'A{header_resp}:'
            f'F{fim_resp}'
        )

    _larguras(
        ws,
        {
            'A': 28,
            'B': 14,
            'C': 14,
            'D': 18,
            'E': 17,
            'F': 16,
            'G': 3,
            'H': 16,
            'I': 16,
            'J': 16,
        },
    )

    ws.freeze_panes = 'A3'

    return ws


# =============================================================================
# ABA CANCELAMENTOS
# =============================================================================
def _montar_cancelamentos(
    wb,
    cancelamentos,
):
    ws = wb.create_sheet(
        'Cancelamentos'
    )

    _sem_grade(ws)

    headers = [
        'Ticket Arion',
        'Designação',
        'Cliente',
        'Cliente Final',
        'ISP',
        'Cidade',
        'UF',
        'Responsável',
        'Fila Atual',
        'Status Atual',
        'Data Criação',
        'Qtd. Pendências',
        'Valor Pendências',
        'Valor Final',
        'Data Cancelamento',
        'Conclusão',
        'Observação Financeira',
    ]

    total_cols = len(
        headers
    )

    _titulo_planilha(
        ws,
        'Cancelamentos',
        '',
        total_cols,
    )

    # Linha 2: agrupamento visual.
    _grupo_colunas(
        ws,
        2,
        1,
        4,
        'IDENTIFICAÇÃO',
        COR_ROXO,
    )

    _grupo_colunas(
        ws,
        2,
        5,
        10,
        'OPERAÇÃO',
        COR_AZUL,
    )

    _grupo_colunas(
        ws,
        2,
        11,
        11,
        'DATAS',
        COR_AMARELO,
    )

    _grupo_colunas(
        ws,
        2,
        12,
        16,
        'FINANCEIRO / ENCERRAMENTO',
        COR_VERDE,
    )

    _grupo_colunas(
        ws,
        2,
        17,
        17,
        'OBS. FINANCEIRA',
        COR_NAVY_2,
    )

    header_row = 3

    for col, header in enumerate(
        headers,
        start=1,
    ):
        ws.cell(
            header_row,
            col,
            header,
        )

    _cabecalho(
        ws,
        header_row,
        1,
        total_cols,
    )

    row = header_row + 1

    for c in cancelamentos:
        validas = (
            _pendencias_validas(
                c
            )
        )

        valor_pendencias = sum(
            _numero(
                p.valor
            )
            for p in validas
        )

        valores = [
            c.ticket_arion,
            c.designacao_circuito,
            c.cliente,
            c.cliente_final,
            c.isp,
            c.cidade,
            c.uf,
            _nome_usuario(
                c.responsavel
            ),
            c.get_fila_display(),
            c.get_status_display(),
            _data_excel(
                c.data_criacao_chamado
            ),
            len(
                validas
            ),
            valor_pendencias,
            _numero(
                c.valor_total_financeiro
            ),
            _data_excel(
                c.data_cancelamento_confirmado
            ),
            _data_excel(
                c.concluido_em
            ),
            c.observacao_financeira,
        ]

        for col, valor in enumerate(
            valores,
            start=1,
        ):
            ws.cell(
                row,
                col,
                valor,
            )

        row += 1

    fim = row - 1

    if fim >= header_row + 1:
        _estilizar_corpo(
            ws,
            header_row + 1,
            fim,
            total_cols,
            wrap_cols={
                3,
                4,
                10,
                17,
            },
            center_cols={
                6,
                7,
                9,
                10,
                11,
                12,
                13,
                14,
                15,
                16,
            },
        )

        _numero_moeda(
            ws,
            [
                'M',
                'N',
            ],
            header_row + 1,
            fim,
        )

        _numero_data(
            ws,
            [
                'K',
                'O',
            ],
            header_row + 1,
            fim,
        )

        _numero_data(
            ws,
            [
                'P',
            ],
            header_row + 1,
            fim,
            com_hora=True,
        )

        # APENAS AutoFilter. Sem Table estruturada.
        ws.auto_filter.ref = (
            f'A{header_row}:'
            f'{get_column_letter(total_cols)}{fim}'
        )

    _larguras(
        ws,
        {
            'A': 18,
            'B': 22,
            'C': 20,
            'D': 24,
            'E': 14,
            'F': 18,
            'G': 8,
            'H': 23,
            'I': 20,
            'J': 25,
            'K': 15,
            'L': 15,
            'M': 18,
            'N': 17,
            'O': 17,
            'P': 19,
            'Q': 38,
        },
    )

    # Congela Ticket/Designação e cabeçalho.
    ws.freeze_panes = 'C4'

    # Ajustes de impressão/visual.
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_view.zoomScale = 85

    return ws


# =============================================================================
# ABA PENDÊNCIAS
# =============================================================================
def _montar_pendencias(
    wb,
    cancelamentos,
):
    ws = wb.create_sheet(
        'Pendências'
    )

    _sem_grade(ws)

    headers = [
        'Ticket Arion',
        'Designação',
        'Cliente Final',
        'Fila do Chamado',
        'Tipo',
        'Descrição',
        'Valor',
        'Status',
        'Data Prevista',
        'Criada em',
        'Criada por',
        'Responsável',
    ]

    total_cols = len(
        headers
    )

    _titulo_planilha(
        ws,
        'Pendências',
        '',
        total_cols,
    )

    _grupo_colunas(
        ws,
        2,
        1,
        4,
        'CHAMADO',
        COR_AZUL,
    )

    _grupo_colunas(
        ws,
        2,
        5,
        8,
        'PENDÊNCIA',
        COR_AMARELO,
    )

    _grupo_colunas(
        ws,
        2,
        9,
        10,
        'DATAS',
        COR_ROXO,
    )

    _grupo_colunas(
        ws,
        2,
        11,
        12,
        'RESPONSÁVEIS',
        COR_VERDE,
    )

    header_row = 3

    for col, header in enumerate(
        headers,
        start=1,
    ):
        ws.cell(
            header_row,
            col,
            header,
        )

    _cabecalho(
        ws,
        header_row,
        1,
        total_cols,
    )

    row = header_row + 1

    for c in cancelamentos:
        for p in c.pendencias.all():
            valores = [
                c.ticket_arion,
                c.designacao_circuito,
                (
                    c.cliente_final
                    or c.cliente
                ),
                c.get_fila_display(),
                p.get_tipo_display(),
                p.descricao,
                _numero(
                    p.valor
                ),
                p.get_status_display(),
                _data_excel(
                    p.data_prevista
                ),
                _data_excel(
                    p.criado_em
                ),
                _nome_usuario(
                    p.criada_por,
                    'Sistema',
                ),
                _nome_usuario(
                    c.responsavel
                ),
            ]

            for col, valor in enumerate(
                valores,
                start=1,
            ):
                ws.cell(
                    row,
                    col,
                    valor,
                )

            row += 1

    fim = row - 1

    if fim >= header_row + 1:
        _estilizar_corpo(
            ws,
            header_row + 1,
            fim,
            total_cols,
            wrap_cols={
                3,
                4,
                5,
                6,
                11,
                12,
            },
            center_cols={
                4,
                7,
                8,
                9,
                10,
            },
        )

        _numero_moeda(
            ws,
            [
                'G',
            ],
            header_row + 1,
            fim,
        )

        _numero_data(
            ws,
            [
                'I',
            ],
            header_row + 1,
            fim,
        )

        _numero_data(
            ws,
            [
                'J',
            ],
            header_row + 1,
            fim,
            com_hora=True,
        )

        # Único AutoFilter, sem Table.
        ws.auto_filter.ref = (
            f'A{header_row}:'
            f'{get_column_letter(total_cols)}{fim}'
        )

        # Destaca status sem quebrar a zebra.
        for linha in range(
            header_row + 1,
            fim + 1,
        ):
            status = str(
                ws.cell(
                    linha,
                    8,
                ).value
                or ''
            ).casefold()

            celula_status = ws.cell(
                linha,
                8,
            )

            if 'resolvida' in status:
                celula_status.fill = _fill(
                    COR_VERDE_CLARO
                )
                celula_status.font = Font(
                    color=COR_VERDE,
                    bold=True,
                    size=9,
                )

            elif 'descartada' in status:
                celula_status.fill = _fill(
                    COR_CINZA_ZEBRA
                )
                celula_status.font = Font(
                    color=COR_CINZA,
                    italic=True,
                    size=9,
                )

            elif 'tratativa' in status:
                celula_status.fill = _fill(
                    COR_AZUL_CLARO
                )
                celula_status.font = Font(
                    color=COR_AZUL,
                    bold=True,
                    size=9,
                )

            else:
                celula_status.fill = _fill(
                    COR_AMARELO_CLARO
                )
                celula_status.font = Font(
                    color=COR_AMARELO,
                    bold=True,
                    size=9,
                )

    _larguras(
        ws,
        {
            'A': 18,
            'B': 22,
            'C': 24,
            'D': 21,
            'E': 23,
            'F': 38,
            'G': 16,
            'H': 17,
            'I': 16,
            'J': 19,
            'K': 22,
            'L': 22,
        },
    )

    ws.freeze_panes = 'B4'
    ws.sheet_view.zoomScale = 90

    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0

    return ws


# =============================================================================
# GERAÇÃO / DOWNLOAD
# =============================================================================
def _gerar_workbook(
    cancelamentos,
    usuario,
):
    wb = Workbook()

    wb.properties.title = (
        'Relatório de Cancelamentos — ARION Broker'
    )

    wb.properties.subject = (
        'Relatório operacional de cancelamentos'
    )

    wb.properties.creator = _nome_usuario(
        usuario,
        'ARION Broker',
    )

    wb.properties.company = (
        'ARION'
    )

    _montar_resumo(
        wb,
        cancelamentos,
        usuario,
    )

    _montar_cancelamentos(
        wb,
        cancelamentos,
    )

    _montar_pendencias(
        wb,
        cancelamentos,
    )

    # Abre sempre no Resumo.
    wb.active = 0

    return wb


@login_required
@require_http_methods(['GET'])
def exportar_relatorio_excel(request):
    """
    Mantém a mesma rota já usada pelo frontend:
    /cancelamentos/relatorio/excel/

    O frontend não precisa ser alterado.
    """
    if Workbook is None:
        return HttpResponse(
            (
                'A biblioteca openpyxl não está instalada. '
                'Execute: pip install openpyxl'
            ),
            status=500,
            content_type='text/plain; charset=utf-8',
        )

    cancelamentos = list(
        _queryset_relatorio(
            request
        )
    )

    workbook = _gerar_workbook(
        cancelamentos,
        request.user,
    )

    # Salva em memória antes de responder.
    # É mais seguro do que escrever diretamente no HttpResponse.
    buffer = BytesIO()

    workbook.save(
        buffer
    )

    buffer.seek(0)

    nome_arquivo = (
        'relatorio_cancelamentos_'
        f'{timezone.localtime():%Y%m%d_%H%M%S}'
        '.xlsx'
    )

    response = HttpResponse(
        buffer.getvalue(),
        content_type=(
            'application/vnd.openxmlformats-officedocument.'
            'spreadsheetml.sheet'
        ),
    )

    response[
        'Content-Disposition'
    ] = (
        f'attachment; filename="{nome_arquivo}"'
    )

    response[
        'Content-Length'
    ] = str(
        buffer.getbuffer().nbytes
    )

    response[
        'Cache-Control'
    ] = 'no-store'

    return response